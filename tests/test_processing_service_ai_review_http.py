from __future__ import annotations

import importlib.util
import io
import json
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

from app.core.config import Settings
from app.schemas.chunk import ChunkOptions
from app.services.processing_service import ProcessingService
from app.storage.repository import JsonRepository


@unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx is not installed")
class ProcessingServiceAIReviewHTTPTests(unittest.TestCase):
    def test_partial_failure_retry_and_reupload_keep_review_contents_without_approval(self) -> None:
        from docx import Document

        requests: list[dict] = []
        failed_once: set[str] = set()
        finding = "제1조의 ‘합성 본문’ 문구를 원문과 대조하세요."

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self) -> None:
                payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                sent = json.loads(payload["messages"][1]["content"])
                item = sent["items"][0]
                requests.append({"path": self.path, "item": item})
                # One batch fails while others persist their useful findings.
                if "제2조" in item["text"] and not failed_once:
                    failed_once.add(item["chunk_id"])
                    self.send_response(503)
                    self.end_headers()
                    self.wfile.write(b'{"error":"synthetic transient failure"}')
                    return
                items = []
                if "제1조" in item["text"]:
                    items.append({
                        "chunk_id": item["chunk_id"], "risk_level": "medium",
                        "issues": [finding], "recommended_human_check": "합성 원문 확인",
                    })
                review = json.dumps({"items": items}, ensure_ascii=False)
                body = json.dumps({
                    "id": "synthetic-request",
                    "choices": [{"finish_reason": "stop", "message": {"content": f"```json\n{review}\n```"}}],
                    "usage": {"prompt_tokens": 20, "completion_tokens": 10, "total_tokens": 30},
                }, ensure_ascii=False).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, format: str, *args: object) -> None:
                pass

        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(server.server_close)
        self.addCleanup(thread.join, 5)
        self.addCleanup(server.shutdown)
        source = Document()
        source.add_paragraph("합성 검토 규정")
        for number in range(1, 4):
            source.add_paragraph(f"제{number}조(시험{number}) 합성 본문 {number}을 확인한다.")
        buffer = io.BytesIO()
        source.save(buffer)

        with tempfile.TemporaryDirectory() as tmp, patch.dict("os.environ", {"NO_PROXY": "127.0.0.1"}):
            settings = Settings(
                data_dir=Path(tmp), enable_agent_review=True, enable_kordoc_table_parser=False,
                llm_provider="openai-compatible", openai_compatible_api_key="",
                agent_review_api_base_url=f"http://127.0.0.1:{server.server_port}/v1",
                agent_review_model="synthetic-review", agent_review_chunks_per_request=1,
                agent_review_max_attempts=1, agent_review_max_parallel_requests=1,
            )
            repository = JsonRepository(settings)
            service = ProcessingService(settings=settings, repository=repository)
            document = service.documents.upload("synthetic.docx", buffer.getvalue(), tenant_id="tenant-a")
            options = ChunkOptions(enable_agent_review=True)
            first_job = service.process(document.document_id, options)
            first_run = repository.latest_completed_run(document.document_id)
            first_review = first_run.stats["agent_review"]
            self.assertEqual("completed", first_job.status)
            self.assertEqual("provider_partial_batches_failed", first_review["skip_reason"])
            self.assertEqual(1, first_review["failed_batch_count"])
            self.assertEqual(list(failed_once), first_review["unreviewed_chunk_ids"])
            calls_after_first = len(requests)
            original_texts = [chunk.text for chunk in repository.get_chunks(document.document_id)]

            retry_job = service.process(document.document_id, options)
            retry_run = repository.latest_completed_run(document.document_id)
            retry_review = retry_run.stats["agent_review"]
            self.assertEqual("completed", retry_job.status)
            self.assertEqual("executed", retry_review["status"])
            self.assertEqual(1, len(requests) - calls_after_first)
            self.assertEqual(1, retry_review["reviewed_chunk_count"])
            self.assertGreater(retry_review["reused_chunk_count"], 0)
            self.assertEqual(0, retry_review["failed_batch_count"])
            chunks = JsonRepository(settings).get_chunks(document.document_id)
            self.assertEqual(original_texts, [chunk.text for chunk in chunks])
            self.assertTrue(any(finding in (chunk.metadata.get("agent_review_findings") or {}).get("issues", []) for chunk in chunks))
            export = json.loads(Path(retry_run.artifacts["ai_review_draft.json"]).read_text(encoding="utf-8"))
            self.assertEqual(retry_review["reviewed_chunk_ids"], export["reviewed_chunk_ids"])
            self.assertGreater(export["reused_finding_count"], 0)
            self.assertEqual([], repository.list_approval_records(document.document_id))
            self.assertEqual("draft", repository.get_document(document.document_id).regulation_status)

            calls_after_retry = len(requests)
            service.process(document.document_id, options)
            self.assertEqual(calls_after_retry, len(requests))
            reupload = service.documents.upload("synthetic.docx", buffer.getvalue(), tenant_id="tenant-a")
            service.process(reupload.document_id, options)
            self.assertEqual(calls_after_retry, len(requests))
            reused_review = repository.latest_completed_run(reupload.document_id).stats["agent_review"]
            self.assertEqual("review_candidates_cached", reused_review["skip_reason"])
            self.assertGreater(reused_review["reused_finding_count"], 0)
            other_tenant = service.documents.upload("synthetic.docx", buffer.getvalue(), tenant_id="tenant-b")
            service.process(other_tenant.document_id, options)
            self.assertGreater(len(requests), calls_after_retry)
            self.assertTrue(all(request["path"] == "/v1/chat/completions" for request in requests))


if __name__ == "__main__":
    unittest.main()
