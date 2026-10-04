from __future__ import annotations

import io
import json
import socket
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import Mock, patch

try:
    from streamlit.testing.v1 import AppTest
except ImportError:  # pragma: no cover - optional operator UI dependency
    AppTest = None

from app.core import config as config_module
from app.core.config import Settings
from app.core.tenant_access import institution_storage_dir
from app.schemas.chunk import Chunk
from app.storage.repository import JsonRepository
from frontend import streamlit_app
from tests.test_streamlit_approval_app import _seed_app_institution_context


APP_PATH = Path(__file__).resolve().parents[1] / "frontend" / "streamlit_app.py"
FINDING = "합성 조항의 경계를 원문과 대조하세요."


class AIReviewStatusTests(unittest.TestCase):
    def test_partial_and_failed_statuses_never_claim_completion_or_expose_errors(self) -> None:
        for status, extra, expected in (
            ("executed", {"failed_batch_count": 1, "unreviewed_chunk_ids": ["b"]}, "일부 완료"),
            ("provider_execution_failed", {}, "실패"),
            ("provider_execution_blocked", {}, "전송 차단"),
            ("api_configuration_needed", {"skip_reason": "agent_review_api_disabled"}, "실행 안 됨"),
        ):
            with self.subTest(status=status):
                tag, message, complete = streamlit_app._ai_review_status_text({
                    "status": status, "request_enabled": True,
                    "provider_error": "private-provider-error", **extra,
                })
                self.assertIn(expected, tag)
                self.assertFalse(complete)
                self.assertNotIn("private-provider-error", message)

    def test_explicit_coverage_does_not_count_unreviewed_selected_chunks(self) -> None:
        summary = {
            "status": "executed", "selected_candidates": [{"chunk_id": "a"}, {"chunk_id": "b"}],
            "reviewed_chunk_ids": ["a"], "unreviewed_chunk_ids": ["b"],
            "reused_candidates": [{"chunk_id": "cached"}],
        }
        self.assertEqual({"a", "cached"}, streamlit_app._agent_review_reviewed_chunk_ids(summary))

    def test_new_backend_failure_reasons_have_actionable_korean_guidance(self) -> None:
        for reason in (
            "provider_response_invalid_review_schema", "provider_response_unknown_chunk",
            "provider_response_duplicate_chunk", "provider_response_truncated",
        ):
            with self.subTest(reason=reason):
                text = streamlit_app._ai_review_retry_guidance({"skip_reason": reason})
                self.assertIn("다시 전처리", text)
                self.assertNotIn(reason, text)
                self.assertIn("길이 제한" if reason.endswith("truncated") else "JSON", text)

    def test_same_title_regulations_with_distinct_numbers_stay_separate(self) -> None:
        chunks = [
            Chunk(chunk_id=f"chunk-{number}", document_id="synthetic", chunk_type="article", text="합성",
                  metadata={"regulation_title": "합성 규정", "regulation_no": str(number)})
            for number in (1, 2)
        ]
        rows = streamlit_app._ai_review_work_rows(chunks, {})
        self.assertNotEqual(rows[0]["규정"], rows[1]["규정"])


@unittest.skipIf(AppTest is None, "Streamlit AppTest is unavailable")
class AIReviewUIFixture(unittest.TestCase):
    """Drive real UI + parser + HTTP transport with synthetic, local-only data.

    Holds no tests so other modules can reuse the fixture without rerunning these.
    """

    def setUp(self) -> None:
        self.root = Path(self.enterContext(tempfile.TemporaryDirectory()))
        self.requests: list[dict] = []
        self.fail_first = False
        self.clean = False
        self.synthetic_bytes: bytes | None = None
        case = self

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self) -> None:
                payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                items = json.loads(payload["messages"][1]["content"])["items"]
                case.requests.append({"path": self.path, "items": items})
                if case.fail_first and len(case.requests) == 1:
                    self.send_response(503)
                    self.end_headers()
                    self.wfile.write(b'{"error":"synthetic unavailable"}')
                    return
                findings = [] if case.clean else [{
                    "chunk_id": item["chunk_id"], "risk_level": "medium",
                    "issues": [FINDING], "recommended_human_check": "합성 원문 확인",
                } for item in items]
                review = json.dumps({"items": findings}, ensure_ascii=False)
                body = json.dumps({
                    "id": "synthetic-ui-review",
                    "choices": [{"finish_reason": "stop", "message": {"content": review}}],
                    "usage": {"prompt_tokens": 10, "completion_tokens": 10, "total_tokens": 20},
                }, ensure_ascii=False).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, format: str, *args: object) -> None:
                pass

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(thread.join, 5)
        self.addCleanup(self.server.shutdown)
        self.enterContext(patch.dict("os.environ", {"NO_PROXY": "127.0.0.1"}))
        self.endpoint = f"http://127.0.0.1:{self.server.server_port}/v1"
        self.settings = Settings(
            app_env="test", data_dir=self.root / "data", artifact_root=self.root,
            institution_profiles_path=str(self.root / "profiles.json"), quality_profiles_path="",
            api_auth_required=False, tenant_storage_isolation=False, api_default_tenant_id="default",
            enable_agent_review=False, llm_provider="openai", agent_review_model="",
            agent_review_api_base_url="https://api.openai.com", agent_review_max_attempts=1,
            agent_review_chunks_per_request=1, agent_review_max_parallel_requests=1,
            agent_review_max_chunks_per_document=0, agent_review_max_input_tokens_per_document=0,
            local_structure_review_enabled=False, enable_kordoc_table_parser=False,
            kordoc_table_command="", pdf_ocr_backend="", rag_llm_backend="extractive",
            openai_api_key="", openai_compatible_api_key="", azure_openai_api_key="",
            azure_openai_endpoint="", anthropic_api_key="",
        )
        self.enterContext(patch.object(config_module, "_runtime_overrides", {}))
        self.enterContext(patch.object(config_module, "_base_settings", return_value=self.settings))
        # Allow only our local provider and Windows asyncio's internal socketpair.
        pair_scope = threading.local()
        original_pair = socket.socketpair
        original_connect = socket.socket.connect
        self.external_network = Mock(side_effect=AssertionError("External networking is forbidden"))

        def socketpair(*args, **kwargs):
            pair_scope.active = True
            try:
                return original_pair(*args, **kwargs)
            finally:
                pair_scope.active = False

        def connect(sock, address):
            if getattr(pair_scope, "active", False) or address == ("127.0.0.1", self.server.server_port):
                return original_connect(sock, address)
            return self.external_network(address)

        self.enterContext(patch.object(socket, "socketpair", socketpair))
        self.enterContext(patch.object(socket.socket, "connect", connect))
        self.addCleanup(self.external_network.assert_not_called)
        self.app = AppTest.from_file(str(APP_PATH), default_timeout=35)
        self.app.session_state["ai_connection_overrides"] = vars(self.settings).copy()
        self.app.session_state["beginner_guide_enabled"] = False
        self.app.session_state["beginner_guide_choice_made"] = True
        self.app.session_state["nav_page"] = streamlit_app.NAV_PREPROCESS
        _seed_app_institution_context(self.app)

    def _run(self) -> None:
        self.app.run()
        self.assertFalse(self.app.exception)

    def _configure(self) -> None:
        self._run()
        self.app.toggle(key="sidebar-ai-review-enabled").set_value(True).run()
        self.app.selectbox(key="sidebar-ai-review-provider").set_value("openai-compatible").run()
        self.app.text_input(key="sidebar-ai-review-model-direct-openai-compatible").input("synthetic-review")
        self.app.text_input(key="sidebar-ai-review-base-url-openai-compatible").input(self.endpoint)
        self.app.button(key="sidebar-ai-review-save").click().run()
        self.assertFalse(self.app.exception)
        saved = self.app.session_state["ai_connection_overrides"]
        self.assertTrue(saved["enable_agent_review"])
        self.assertEqual(self.settings.data_dir, saved["data_dir"])
        self.assertEqual(1, saved["agent_review_max_attempts"])
        self.assertEqual("openai-compatible", self.app.selectbox(key="sidebar-ai-review-provider").value)
        self.assertTrue(self.app.toggle(key="sidebar-ai-review-enabled").value)

    def _process(self) -> str:
        from docx import Document

        if self.synthetic_bytes is None:
            source = Document()
            source.add_paragraph("합성 AI 검수 규정")
            for number in range(1, 4):
                source.add_paragraph(f"제{number}조(시험{number}) 합성 본문 {number}을 확인한다.")
            buffer = io.BytesIO()
            source.save(buffer)
            self.synthetic_bytes = buffer.getvalue()
        pending = institution_storage_dir(self.settings.data_dir / "pending_uploads", "test-profile", create=True)
        # Reuse byte-for-byte identical input so the UI run-cache path is tested.
        (pending / "synthetic-ai-review.docx").write_bytes(self.synthetic_bytes)
        self._run()
        self.app.button(key="pending-upload-select-all").click().run()
        self.assertFalse(self.app.button(key="preprocess-start").disabled)
        self.app.button(key="preprocess-start").click().run()
        self.assertFalse(self.app.exception)
        return self.app.session_state["document_id"]

    def _open_results(self, document_id: str) -> None:
        self.app.session_state["nav_page"] = streamlit_app.NAV_RESULTS
        self.app.session_state["workflow_opened_document_id"] = document_id
        self._run()

    def _work_table(self):
        return next(frame.value for frame in self.app.dataframe if "AI 작업 상태" in frame.value.columns)


class AIReviewUIWorkflowTests(AIReviewUIFixture):

    def test_sidebar_configuration_partial_findings_and_reupload_retry(self) -> None:
        self.fail_first = True
        self._configure()
        document_id = self._process()
        repository = JsonRepository(self.settings)
        first_summary = repository.list_runs(document_id)[-1].stats["agent_review"]
        self.assertEqual(1, first_summary["failed_batch_count"])
        self.assertTrue(any("AI 검수 일부 완료" in str(item.value) for item in self.app.text))
        self._open_results(document_id)
        table = self._work_table()
        self.assertIn(FINDING, "\n".join(table["검수 의견"]))
        self.assertIn("미완료 · 결과 없음", list(table["AI 작업 상태"]))
        self.assertTrue(any("AI 검수 일부 완료" in item.value for item in self.app.info))
        calls_before_retry = len(self.requests)
        self.app.button(key=f"results-ai-work-{document_id}-retry").click().run()
        if "_nav_target" in self.app.session_state:
            self._run()
        self.assertEqual(streamlit_app.NAV_PREPROCESS, self.app.session_state["nav_page"])
        second_id = self._process()
        second_summary = repository.list_runs(second_id)[-1].stats["agent_review"]
        self.assertEqual(0, second_summary["failed_batch_count"])
        self.assertEqual(1, len(self.requests) - calls_before_retry)
        self._open_results(second_id)
        self.assertNotIn("미완료 · 결과 없음", list(self._work_table()["AI 작업 상태"]))
        self.assertIn(FINDING, "\n".join(self._work_table()["검수 의견"]))
        for current_id in (document_id, second_id):
            self.assertEqual([], repository.list_approval_records(current_id))
            self.assertTrue(all(chunk.approval_status != "approved" for chunk in repository.get_chunks(current_id)))
        self.assertFalse(list(self.settings.data_dir.rglob("approved_vectors.jsonl")))

    def test_clean_response_is_visible_as_reviewed_without_findings(self) -> None:
        self.clean = True
        self._configure()
        document_id = self._process()
        self._open_results(document_id)
        table = self._work_table()
        self.assertTrue(all(value == "검수 완료 · 지적 없음" for value in table["AI 작업 상태"]))
        self.assertTrue(any("AI 검수 초안 완료" in item.value for item in self.app.success))
        self.assertTrue(self.requests)

    def test_invalid_configuration_does_not_enable_review(self) -> None:
        self._run()
        self.app.toggle(key="sidebar-ai-review-enabled").set_value(True).run()
        self.app.button(key="sidebar-ai-review-save").click().run()
        self.assertFalse(self.app.exception)
        self.assertFalse(self.app.session_state["ai_connection_overrides"]["enable_agent_review"])
        self.assertTrue(self.app.error)
        self.assertEqual([], self.requests)

    def test_regulation_selector_keeps_findings_and_failed_rows_separate(self) -> None:
        chunks = [
            Chunk(chunk_id="a", document_id="synthetic", chunk_type="article", text="합성 A",
                  metadata={"regulation_title": "합성 인사규정", "agent_review_findings": {"issues": [FINDING]}}),
            Chunk(chunk_id="b", document_id="synthetic", chunk_type="article", text="합성 B",
                  metadata={"regulation_title": "합성 복무규정"}),
        ]
        app = AppTest.from_string(
            "from frontend.streamlit_app import _render_ai_review_work\n"
            "import streamlit as st\n"
            "_render_ai_review_work(st.session_state['context'], key='fixture')"
        )
        app.session_state["context"] = {
            "chunks": chunks, "agent_review_summary": {
                "status": "executed", "request_enabled": True, "failed_batch_count": 1,
                "selected_candidates": [{"chunk_id": "a"}, {"chunk_id": "b"}],
                "reviewed_chunk_ids": ["a"], "unreviewed_chunk_ids": ["b"],
            },
        }
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual([FINDING], list(app.dataframe[0].value["검수 의견"]))
        app.selectbox(key="fixture-regulation").set_value("합성 복무규정").run()
        self.assertEqual(["미완료 · 결과 없음"], list(app.dataframe[0].value["AI 작업 상태"]))
        self.assertEqual([""], list(app.dataframe[0].value["검수 의견"]))


if __name__ == "__main__":
    unittest.main()
